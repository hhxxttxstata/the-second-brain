# -*- coding: utf-8 -*-
"""生成秋招投递管理 Excel 表（数据来自 QQ 邮箱近两周公司邮件）"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

rows = [
    # 公司, 岗位, 来源, 状态, 最近动态(邮件日期), 备注
    ("北京清微智能科技", "AI软件测试开发工程师", "官网招聘系统(iTalent/北森)", "已投递（简历已收）", "2026-08-25", "感谢投递邮件 UID906"),
    ("北京清微智能科技", "AI软件开发工程师（C/C++）", "官网招聘系统(iTalent/北森)", "已投递（简历已收）", "2026-08-25", "感谢投递邮件 UID905"),
    ("摩尔线程", "AI测试工程师(J10987)", "官网招聘系统(iTalent/北森)", "已投递（简历已收）", "2026-08-25", "感谢投递邮件 UID904"),
    ("摩尔线程", "Agent系统研发工程师(J10988)", "官网招聘系统(iTalent/北森)", "已投递（简历已收）", "2026-08-25", "感谢投递邮件 UID903"),
    ("途牛旅游网", "【2027届校招】AI应用开发工程师(J12964)", "官网招聘系统(北森)", "已投递（简历已收）", "2026-08-25", "感谢投递邮件 UID902"),
    ("卡尔动力", "27届校招算法岗（自动驾驶研究员/感知/控制/决策规划/定位标定）", "官网(飞书招聘)HR邀请", "HR联系中（意向收集）", "2026-08-25", "此前投过实习岗，HR发意向收集链接 UID901"),
    ("4399游戏", "【2027校招】AI Agent技术开发岗", "官网(4399校招)", "未通过（感谢信）", "2026-08-24", "8/20发笔试邀请(UID875)，8/24感谢信称暂不匹配(UID898)；8/19曾提醒完善简历(UID871)"),
    ("冠宇集团（珠海冠宇电池）", "校招岗位（邮件未写明具体岗位）", "官网招聘系统(北森)", "测评中（待完成）", "2026-08-24", "简历投递成功(UID897)+在线测评邀请，9/8前完成(UID896)"),
    ("杭州广立微电子", "AI应用工程师（27届）-杭州", "官网(Moka)", "已投递（简历已收）", "2026-08-24", "感谢投递邮件 UID895"),
    ("中联重科", "机器人研发工程师（具身智能与AI算法）", "官网(Wintalent招聘系统)", "测评中（待完成）", "2026-08-24", "人才测评通知，8/27前完成 UID894"),
    ("顺丰科技", "AI开发工程师", "官网（赛码笔试系统）", "笔试邀请（8/23开放）", "2026-08-23", "专业笔试邀请函 UID892/887/885（同一次笔试三次提醒）"),
    ("顺丰集团", "校招岗位（邮件未写明具体岗位）", "官网（顺丰招聘）", "测评中（待完成）", "2026-08-18", "校园招聘测评邀请，24小时内完成 UID868"),
    ("英科", "未明确（邮件未写明岗位）", "官网(飞书招聘)", "未通过（感谢信）", "2026-08-22", "简历筛选未通过，已录入人才库 UID888"),
    ("小鹏汽车", "【27届校招】AI前瞻培训生", "官网(飞书招聘)", "AI测评中（待完成）", "2026-08-21", "AI测评面试邀请，72小时内完成 UID884"),
    ("海能达", "AI算法工程师（Agent开发方向）（27届校招）", "官网(Moka)", "已投递（待补充简历信息）", "2026-08-20", "需更新简历及面试站点信息 UID879"),
    ("华诺星空", "ai算法工程师", "官网(Moka)/HR邮箱", "已投递（简历已收）", "2026-08-19", "感谢投递邮件 UID872；HR另发2027届校招简章 UID891"),
    ("蔚来NIO", "AGI超星计划-面向销售的Agentic强化学习研究", "官网(飞书招聘)", "已投递（简历已收）", "2026-08-17", "感谢投递邮件 UID866"),
    ("中望软件", "AI算法工程师（2027届）-ZWCAD", "官网(JoinUs招聘系统)", "已投递（问卷待填）", "2026-08-17", "感谢投递邮件 UID864；另有校招生问卷填写邀请 UID865"),
    ("恒生电子", "AI应用开发工程师(J10355)", "官网招聘系统(北森)", "已投递（简历已收）", "2026-08-17", "感谢投递邮件 UID862"),
    ("基恩士（中国）", "销售工程师（2027届秋招）", "官网/智联招聘", "AI面试中（待完成）", "2026-08-25", "多次AI面试邀请 UID900/893/883/863/860/853，8/28前完成；岗位信息来自智联推荐邮件 UID907"),
]

wb = Workbook()
ws = wb.active
ws.title = "秋招投递管理"

headers = ["投递公司", "投递岗位", "投递来源", "投递状态", "最近动态日期", "备注"]

# 样式
header_fill = PatternFill("solid", fgColor="4472C4")
header_font = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
body_font = Font(name="微软雅黑", size=10)
center = Alignment(horizontal="center", vertical="center", wrap_text=True)
left = Alignment(horizontal="left", vertical="center", wrap_text=True)
thin = Side(style="thin", color="B0B0B0")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

status_colors = {
    "未通过": "FCE4E4",
    "测评中": "FFF2CC",
    "笔试": "FFF2CC",
    "面试": "E2EFDA",
    "HR联系": "DDEBF7",
}

for col, h in enumerate(headers, 1):
    c = ws.cell(row=1, column=col, value=h)
    c.fill = header_fill
    c.font = header_font
    c.alignment = center
    c.border = border

for r, row in enumerate(rows, 2):
    for col, v in enumerate(row, 1):
        c = ws.cell(row=r, column=col, value=v)
        c.font = body_font
        c.border = border
        c.alignment = center if col in (1, 4, 5) else left
        # 状态列着色
        if col == 4:
            for key, color in status_colors.items():
                if key in str(v):
                    c.fill = PatternFill("solid", fgColor=color)
                    break

widths = [20, 42, 26, 22, 14, 60]
for i, w in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w
ws.row_dimensions[1].height = 22
for r in range(2, len(rows) + 2):
    ws.row_dimensions[r].height = 30

ws.freeze_panes = "A2"

out = r"D:\MyAgent\秋招投递管理表.xlsx"
wb.save(out)
print(f"saved: {out}  (rows: {len(rows)})")
